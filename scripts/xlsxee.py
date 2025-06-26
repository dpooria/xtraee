#!/usr/bin/env python
import argparse
import logging
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment

MOLKEY = "Molecules"
IRRKEY = "Irrep"


def load_xlsx(xlsx: str, noclean: bool = False) -> pd.DataFrame:
    wb = load_workbook(xlsx)
    if IRRKEY not in wb.sheetnames:
        logging.error(f"{IRRKEY} sheet not found")
        return 1
    irrep_sheet = wb[IRRKEY]
    for i, row in enumerate(irrep_sheet.iter_rows(values_only=True)):
        if MOLKEY in row:
            break
    else:
        logging.error("Molecule column not found")
        return 1
    if noclean:
        df = pd.DataFrame(irrep_sheet.values, columns=row)
        return df
    # clean up the excel file
    df = (
        pd.DataFrame(
            irrep_sheet.values,
            columns=row,
        )
        .iloc[i + 1 :]  # noqa
        .dropna(how="all", axis=1)
        .dropna(how="all", axis=0)
    )
    molidx = df.columns.get_loc(MOLKEY)
    df = df[df.columns[molidx:].dropna()].reset_index(drop=True)
    # flatten ccsd and cis labels
    data = []
    cols = [MOLKEY]
    for c in df.columns:
        if c == MOLKEY or c is None:
            continue
        cols.append(f"{c}_CCSD")
    for c in df.columns:
        if c == MOLKEY or c is None:
            continue
        cols.append(f"{c}_CIS")
    buf = []
    for i, row in df.iterrows():
        if row[MOLKEY] is None:
            buf.extend(row[1:])
            data.append(dict(zip(cols, buf)))
            buf.clear()
        else:
            buf.append(row[MOLKEY])
            buf.extend(row[1:])

    return pd.DataFrame(data)


CCSD_ST = "A_A_compare_ccsd.csv"
CIS_ST = "singlet_triplet_compare_cis.csv"
CCSDCIS_SS = "CIS_singlet_CCSD_A_compare_ccsd_vs_cis.csv"
CCSDCIS_TT = "CIS_triplet_CCSD_A_compare_ccsd_vs_cis.csv"


def rootdirs(root: Path, molkey: pd.Series) -> list[Path]:
    def getroot(root: Path, mol: str) -> Path:
        transform = [
            lambda x: x,
            lambda x: x.strip().strip("*"),
            lambda x: x.lower().strip().strip("*"),
            lambda x: x.upper().strip().strip("*"),
            lambda x: x.capitalize().strip().strip("*"),
            lambda x: x.replace("-", "_"),
            lambda x: x.replace("_", "-"),
            lambda x: x.replace("-", "_").lower(),
            lambda x: x.replace("_", "-").lower(),
            lambda x: x.replace("-", "_").upper(),
            lambda x: x.replace("_", "-").upper(),
        ]
        for tr in transform:
            if (p := root / tr(mol)).is_dir():
                break
        return p

    paths = []
    for mol in molkey:
        p = getroot(root, mol)
        if p.is_dir():
            paths.append(p)
        else:
            logging.error(
                f"Root directory for molecule {mol} not found, ignoring the molecule"
            )
    for p in paths:
        for c in (CCSD_ST, CIS_ST, CCSDCIS_SS, CCSDCIS_TT):
            if not (p / c).is_file():
                logging.error(f"File {p / c} not found, ignoring the molecule")
                paths.remove(p)
                continue
    return paths


def intable(x: str) -> int | None:
    if x is None:
        return x
    try:
        return int(x)
    except ValueError:
        try:
            return int(x.strip("*"))
        except ValueError:
            return None


def singlet_triplet(
    irrep: pd.DataFrame,
    paths: list[Path],
    file: str,
    suffix: str,
) -> pd.DataFrame:
    singletcol = [c for c in irrep.columns if c.endswith(suffix) and c.startswith("S")]

    if suffix == "_CCSD":
        irr = "A"
    else:
        irr = ""
    result = []
    for (_, row), path in zip(irrep.iterrows(), paths):
        data = pd.read_csv(path / file, index_col=0)
        singletind = row[singletcol].apply(intable).dropna().astype(int)
        loc = data.loc[[f"singlet-{i}/{irr}" for i in singletind]].max(axis=1)
        loc.index = map(lambda x: f"{x.removesuffix(suffix)}-Tn", singletind.index)
        d = {MOLKEY: row[MOLKEY]}
        d.update({k: v for k, v in loc.items()})
        for k in row[singletcol].index:
            t = f"{k.removesuffix(suffix)}-Tn"
            if t not in loc.index:
                d[t] = None
        result.append(d)
    return pd.DataFrame(result)


def mvm(
    irrep: pd.DataFrame,
    paths: list[Path],
    file: str,
    suffix_m: tuple[str, str],
    ee_m: tuple[str, str],
    irr_m: tuple[str, str],
) -> pd.DataFrame:
    """method vs method comparison"""
    assert len(suffix_m) == len(irr_m) == len(ee_m) == 2
    m0col = [
        c for c in irrep.columns if c.endswith(suffix_m[0]) and c.startswith(ee_m[0])
    ]
    eemap = {"S": "singlet", "T": "triplet"}
    result = []
    for (_, row), path in zip(irrep.iterrows(), paths):
        data = pd.read_csv(path / file, index_col=0)
        m0ind = row[m0col].apply(intable).dropna().astype(int)
        loc = data[[f"{eemap[ee_m[0]]}-{i}/{irr_m[0]}" for i in m0ind]].max(axis=0)
        loc.index = map(
            lambda x: f"{x.removesuffix(suffix_m[0])}-{ee_m[1]}n", m0ind.index
        )
        d = {MOLKEY: row[MOLKEY]}
        d.update({k: v for k, v in loc.items()})
        for k in row[m0col].index:
            t = f"{k.removesuffix(suffix_m[0])}-{ee_m[1]}n"
            if t not in loc.index:
                d[t] = None
        result.append(d)
    return pd.DataFrame(result)




def parseargs() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Gather the resulting `xtraee compareall` data of different molecules into a single xlsx file"
    )
    parser.add_argument(
        "xlsx",
        type=str,
        help="Input xlsx file containing the Irrep sheet with the Molecules and their respective Singlet and Triplet labels",
    )
    parser.add_argument(
        "--root",
        type=str,
        default=".",
        help="Root directory where the data is stored (default: current directory)",
    )
    parser.add_argument(
        "--output",
        type=str,
        default="xlsxee.xlsx",
        help="Output file",
    )
    parser.add_argument("--silent", action="store_true", help="silent mode")
    parser.add_argument(
        "--noclean",
        action="store_true",
        help="do not clean the input xlsx file before processing",
    )
    return parser.parse_args()


def main() -> int:
    args = parseargs()
    if not args.silent:
        setup_logging()
    logging.info(f"Read data from {args.xlsx}")
    irrep = load_xlsx(args.xlsx, args.noclean)
    logging.info(f"Read {len(irrep)} molecules from {args.root}")
    paths = rootdirs(Path(args.root), irrep[MOLKEY])
    logging.info(f"Found {len(paths)} root directories")
    logging.info("Read singlet-triplet comparison data")
    acc1 = singlet_triplet(irrep, paths, CCSD_ST, "_CCSD")
    logging.info("Read CIS singlet vs CCSD singlet comparison")
    acc2 = mvm(irrep, paths, CCSDCIS_SS, ("_CCSD", "_CIS"), ("S", "S"), ("A", ""))
    logging.info("Read CIS singlet vs triplet comparison")
    acc3 = singlet_triplet(irrep, paths, CIS_ST, "_CIS")
    logging.info("Read CIS triplet vs CCSD triplet comparison")
    acc4 = mvm(irrep, paths, CCSDCIS_TT, ("_CCSD", "_CIS"), ("T", "T"), ("A", ""))
    logging.info("Write data to xlsx file")
    with pd.ExcelWriter(args.output, mode="w") as writer:
        toexcel_args = dict(startcol=2, startrow=2, index=False)
        irrep.to_excel(writer, sheet_name="Irrep", **toexcel_args)
        acc1.to_excel(writer, sheet_name="Accuracy1", **toexcel_args)
        acc2.to_excel(writer, sheet_name="Accuracy2", **toexcel_args)
        acc3.to_excel(writer, sheet_name="Accuracy3", **toexcel_args)
        acc4.to_excel(writer, sheet_name="Accuracy4", **toexcel_args)
        title = [
            "Irrep",
            "CCSD singlets vs triplets",
            "CCSD singlets vs CIS singlets",
            "CIS singlets vs triplets",
            "CCSD triplets vs CIS triplets",
        ]
        for i, sheet in enumerate(writer.sheets):
            ws = writer.sheets[sheet]
            # write title
            ws.cell(row=1, column=1, value=title[i])
            for i in range(2, ws.max_row + 1):
                ws.row_dimensions[i].height = 20
                for j in range(1, ws.max_column + 1):
                    ws.cell(row=i, column=j).alignment = Alignment(horizontal="center")
            # change the default precision
            if sheet != "Irrep":
                for col in ws.columns:
                    for cell in col:
                        if cell.data_type == "n":
                            cell.number_format = "0.000"
    logging.info(f"Write data to {args.output}")
    logging.info("Done")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
